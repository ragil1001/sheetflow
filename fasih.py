import json
import random
import re
import sys
from contextlib import suppress
from decimal import Decimal, InvalidOperation
from math import ceil
from pathlib import Path
from time import monotonic
from urllib.parse import quote, urlsplit
from camoufox.fingerprints import from_fpgen, generate_fingerprint
from camoufox.sync_api import Camoufox
from camoufox.utils import get_screen_cons
from openpyxl import load_workbook
from playwright.sync_api import Error, expect

TIMEOUT = 30_000
TIMEOUT_CARI_USAHA = 5_000
PESAN_BOT = re.compile(r"\bBOT-\d+\b|\bbot\s+detected\b|(?:terdeteksi|teridentifikasi)\s+(?:sebagai\s+)?(?:bot|robot)\b|koneksi\s+anda\s+sebagai\s+bot\b|(?:kami\s+mendeteksi|sistem\s+kami\s+mendeteksi)\s+(?:perilaku\s+yang\s+tidak\s+wajar|koneksi)|perilaku\s+yang\s+tidak\s+wajar\s+pada\s+koneksi", re.I)

FIELDS = {
    "tk_laki_laki": "Pekerja laki-laki",
    "tk_perempuan": "Pekerja perempuan",
    "tk_dibayar": "Pekerja dibayar",
    "tk_tidak_dibayar": "Pekerja tidak dibayar",
    "upah_gaji": "Total upah dan gaji, serta jaminan sosial pegawai",
    "perbaikan_produksi": "Biaya produksi",
    "perbaikan_biaya_pembelian": "Biaya pembelian barang yang terjual",
    "perbaikan_operasional": "Biaya operasional",
    "perbaikan_nonoperasional": "Biaya non-operasional",
    "perbaikan_pendapatan": "Nilai pendapatan barang dan jasa",
}
SELECTORS = {
    "revoke": 'button:visible:has(svg.tabler-icon-rotate-clockwise)',
    "konfirmasi_revoke": '[role="alertdialog"]:visible:has-text("Revoke Assignment?") button:text-is("Konfirmasi")',
    "edit": 'button:visible:has(svg.tabler-icon-edit)',
    "blok_ii": '.fasih-form-sidebar:visible > div[title="SE2026 - L BLOK II"]:visible',
    "kartu_usaha": '#se2026_nested [data-nested-view="true"]:visible',
    "catatan": '.fasih-form-sidebar:visible > div[title="CATATAN"]:visible',
    "switch": '#cek_anomali_button input[role="switch"]',
    "kontrol_switch": '#cek_anomali_button div[id$="-control"]',
    "modal": r'div.tw\:border\!.tw\:relative.tw\:flex.tw\:w-full.tw\:flex-col.tw\:gap-4.tw\:bg-background.tw\:p-6.tw\:pt-4.tw\:shadow-lg\!.tw\:sm\:rounded-lg:visible',
}

def pola_usaha(target):
    target = " ".join(target.split())
    parts = re.fullmatch(r"(.+?)\s*(?:\(([^()<>]+)\)|<([^()<>]+)>?)", target)
    text = re.escape(target).replace(r"\ ", r"\s+")
    if parts:
        owner = re.escape((parts[2] or parts[3]).strip()).replace(r"\ ", r"\s+")
        text = re.escape(parts[1].strip()).replace(r"\ ", r"\s+") + r"(?:[\s.,/\-]+[^()<>]*)?\s*(?:\(\s*" + owner + r"\s*\)|<\s*" + owner + r"\s*>?)"
    return re.compile(r"^\s*" + text + r"\s*$", re.I)

def inputan(page, label):
    text = r"[\W_]+".join(re.escape(w) for w in re.findall(r"\w+", label))
    if label == FIELDS["perbaikan_pendapatan"]:
        text = r"Nilai\W+(?:(?:produksi|pendapatan|penjualan)\W+){1,3}barang\W+dan\W+jasa"
    prefix = r"(?:(?!\b(?:total|rincian|jumlah|subtotal|petunjuk|contoh|catatan|keterangan|akumulasi)\b)[\s\S])*?"
    pattern = re.compile(r"^" + prefix + r"\b(?:" + text + r")\b[\s\S]*$", re.I)
    inputs = 'input[not(@type) or @type="text" or @type="number" or @type="tel"]'
    rows = page.get_by_text(pattern).filter(visible=True).locator(
        "xpath=ancestor-or-self::div[contains(concat(' ', normalize-space(@class), ' '), ' tw:rounded-lg ') "
        "and contains(concat(' ', normalize-space(@class), ' '), ' tw:gap-2 ')][1]")
    field = page.get_by_label(pattern).and_(page.locator("xpath=//" + inputs))
    return field.or_(rows.locator(f"xpath=self::*[count(.//{inputs})=1]//{inputs}")).filter(visible=True)

def buka_usaha(page, target):
    target = " ".join(target.split())
    cards = page.locator(SELECTORS["kartu_usaha"])
    daftar = page.locator("#se2026_nested")
    kosong = daftar.get_by_text("Tidak ada data", exact=True).filter(visible=True)
    expect(cards.first.or_(kosong).first, "Daftar kartu usaha BLOK II tidak muncul").to_be_visible(timeout=TIMEOUT)
    if kosong.is_visible() and not cards.count():
        raise ValueError("Tidak ada kartu usaha pada assignment ini: " + target)
    ringkas = re.sub(r"\s*[(<].*$", "", target).strip()
    pola_ringkas = re.compile(r"^\s*" + re.escape(ringkas).replace(r"\ ", r"\s+") + r"\s*(?:[(<][\s\S]*)?$", re.I)
    kandidat = daftar.get_by_text(pola_usaha(target)).or_(daftar.get_by_text(pola_ringkas)).filter(visible=True)
    # ponytail: kartu muncul menandai daftar siap; gunakan indikator selesai jika tersedia.
    terlihat = " | ".join(" ".join(text.split()) for text in cards.all_inner_texts())[:1500]
    try:
        expect(kandidat.first).to_be_visible(timeout=TIMEOUT_CARI_USAHA)
    except AssertionError as error:
        if page.is_closed():
            raise
        raise ValueError("Kartu usaha tidak cocok: " + target + " | Kartu terlihat: " + terlihat) from error
    pattern = re.compile(r"^\s*" + re.escape(target).replace(r"\ ", r"\s+") + r"\s*$", re.I)
    name = daftar.get_by_text(pattern).filter(visible=True)
    if not name.count():
        pattern = pola_usaha(target)
        name = daftar.get_by_text(pattern).filter(visible=True)
    if not name.count():
        name = daftar.get_by_text(pola_ringkas).filter(visible=True)
    card = name.locator('xpath=ancestor-or-self::*[@data-nested-view="true"][1]')
    if card.count() > 1:
        raise ValueError("Nama usaha cocok dengan beberapa kartu; tidak memilih sembarang kartu")
    expect(card, "Kartu usaha tidak ditemukan: " + target).to_have_count(1, timeout=TIMEOUT_CARI_USAHA)
    name.first.click()

def isi_remarks(page, field, text):
    wadah = field.locator('xpath=ancestor::*[contains(@id, "#")][1]')
    tombol = wadah.locator(r'button.tw\:size-8.tw\:relative.tw\:text-primary.tw\:\@lg\/form\:flex\!.tw\:hidden:visible:has(> svg.tw\:size-4)')
    expect(tombol, "Tombol catatan pada input yang diubah tidak ditemukan/unik").to_have_count(1, timeout=TIMEOUT)
    tombol.click()
    editor = page.get_by_placeholder("Tambah catatan di sini...", exact=True).filter(visible=True)
    editor.fill(text)
    expect(editor).to_have_value(text, timeout=TIMEOUT)
    page.get_by_role("button", name="Simpan", exact=True).click()
    page.get_by_role("button", name="Dismiss", exact=True).click()
    expect(editor).to_be_hidden(timeout=TIMEOUT)

def angka(value):
    if value is None or not str(value).strip():
        return None
    text = "0" if str(value).strip() in ("-", "–", "—") else str(value).strip()
    if isinstance(value, str):
        text = text.replace(".", "").replace(",", ".")
    number = Decimal(text)
    if isinstance(value, bool) or not number.is_finite() or number < 0:
        raise ValueError("Angka tidak valid")
    return number

def simpan(book, path):
    temp = path.with_suffix(".tmp.xlsx")
    book.save(temp)
    temp.replace(path)

def identitas(path):
    if path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        assert isinstance(data, dict) and isinstance(data.get("navigator"), dict), "Fingerprint tersimpan tidak valid"
        assert "Firefox/" in data["navigator"].get("userAgent", ""), "Fingerprint harus Firefox"
        return data
    data = generate_fingerprint(os="windows", screen=get_screen_cons())
    data.setdefault("window", {})["screenY"] = from_fpgen(data)["window.screenY"]
    path.parent.mkdir(exist_ok=True)
    temp = path.with_suffix(".tmp.json")
    temp.write_text(json.dumps(data), encoding="utf-8")
    temp.replace(path)
    return data

def sukses_kirim(page):
    return page.get_by_text(re.compile(r"Assignment\s+Berhasil\s+Disubmit", re.I)).or_(page.get_by_role("button", name="Tetap di Halaman", exact=True)).filter(visible=True).first

def tunggu_kirim(page, tujuan):
    gagal = page.get_by_text(re.compile(r"\bpre[\s_-]?signed[\s_-]*url\b", re.I)).filter(visible=True).first
    sukses = sukses_kirim(page)
    expect(tujuan.filter(visible=True).or_(gagal).or_(sukses).first).to_be_visible(timeout=TIMEOUT)
    if sukses.is_visible():
        return
    assert not gagal.is_visible(), "Kirim gagal: FASIH menampilkan kesalahan presigned URL; hasil kirim belum dipastikan"

def pesan_galat(page):
    modal = page.locator(SELECTORS["modal"])
    tunggu_kirim(page, modal)
    if sukses_kirim(page).is_visible():
        return ""
    tombol = modal.get_by_role("button", name=re.compile(r"^Galat\s+[1-9]\d*\s+Perlu diperbaiki$"))
    rincian = page.locator(SELECTORS["modal"]).filter(has=page.get_by_text(re.compile(r"^Total:\s*\d+\s*/\s*\d+$")))
    if not rincian.count():
        if not tombol.count():
            return ""
        tombol.click(timeout=1000)
    pesan = rincian.locator(r"ul.tw\:text-destructive > li")
    expect(pesan.first.or_(sukses_kirim(page)).first).to_be_visible(timeout=TIMEOUT)
    if sukses_kirim(page).is_visible():
        return ""
    lines = []
    for item in pesan.all():
        wadah = item.locator("xpath=../..")
        label = wadah.locator("[title]")
        judul = label.first.get_attribute("title") if label.count() else " ".join(wadah.locator("xpath=./div").all_inner_texts())
        lines.append((judul.strip() + " — " if judul else "") + item.inner_text().strip())
    text = "\n".join(lines)
    assert text and len(text) <= 32767, "Pesan Galat kosong atau melebihi batas sel Excel"
    return text

def kirim(page, pulih=None):
    sukses = sukses_kirim(page)
    konfirmasi = page.get_by_role("button", name="Konfirmasi", exact=True).filter(visible=True)
    modal = page.locator(SELECTORS["modal"])
    galat = modal.get_by_role("button", name=re.compile(r"^Galat\s+[1-9]\d*\s+Perlu diperbaiki$")).filter(visible=True)
    rincian = modal.filter(has=page.get_by_text(re.compile(r"^Total:\s*\d+\s*/\s*\d+$")))
    awal = page.locator("body").filter(has_not=modal).get_by_role("button", name="Kirim", exact=True).filter(visible=True)
    terakhir, waktu = None, 0
    terlihat, lapor = None, monotonic()
    while not sukses.is_visible():
        if pulih is not None and pulih():
            raise Error("Kirim terputus oleh refresh; periksa ulang form dan hasil kirim.")
        try:
            expect(page.locator("body")).to_be_visible(timeout=TIMEOUT)
            if pulih is not None and pulih():
                continue
            if sukses.is_visible():
                break
            tombol = None
            if rincian.count():
                tahap = "rincian Galat"
                if rincian.locator(r"ul.tw\:text-destructive > li").first.is_visible():
                    detail = pesan_galat(page)
                    if detail:
                        return detail
            elif galat.count():
                tahap, tombol = "buka Galat", galat
            elif konfirmasi.is_visible():
                tahap, tombol = "Konfirmasi Kirim", konfirmasi
            elif modal.count():
                tahap = "menunggu rincian Galat" if terakhir == "buka Galat" else "Kirim pada modal validasi"
                if terakhir != "buka Galat":
                    tombol = modal.last.get_by_role("button", name="Kirim", exact=True)
            else:
                tahap, tombol = "Kirim awal", awal
            if sukses.is_visible():
                break
            status = tahap
            if tombol is not None:
                if not tombol.is_visible():
                    status += ": tombol belum terlihat/loading"
                elif not tombol.is_enabled():
                    status += ": tombol belum aktif/loading"
                # ponytail: tombol aktif menandai request selesai; batasi pengulangan tahap sama menjadi sekali/detik.
                elif tahap != terakhir or monotonic() - waktu >= 1:
                    # Ikat klik ke elemen lama agar handler refresh tidak mengklik Kirim pada form baru.
                    elemen = tombol.element_handle(timeout=1000)
                    try:
                        elemen.click(timeout=1000)
                    finally:
                        with suppress(Error):
                            elemen.dispose()
                    terakhir, waktu = tahap, monotonic()
                    status += ": diklik, menunggu respons"
        except (Error, AssertionError) as error:
            if page.is_closed() or re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                raise
            if isinstance(error, AssertionError) and "timeout" not in str(error).lower():
                raise
            status = "Menunggu tampilan siap: " + str(error).splitlines()[0][:160]
        if sukses.is_visible():
            break
        if status != terlihat or monotonic() - lapor >= 5:
            print("  [KIRIM] " + status, flush=True)
            terlihat, lapor = status, monotonic()
        page.wait_for_timeout(100)
    return ""

def sesi(browser, page, path, simpan=False):
    if simpan:
        tab = json.loads(path.read_text(encoding="utf-8"))["tab"] if path.exists() else {}
        if not page.is_closed() and page.url.startswith("https://fasih-sm.bps.go.id/"):
            tab = page.evaluate("() => ({...sessionStorage})")
        data = {"browser": browser.storage_state(indexed_db=True), "tab": tab}
        temp = path.with_suffix(".tmp.json")
        temp.write_text(json.dumps(data), encoding="utf-8")
        temp.replace(path)
    elif path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        browser.set_storage_state(data["browser"])
        page.add_init_script("if (location.origin === 'https://fasih-sm.bps.go.id' && !sessionStorage.length) "
                             "for (const [k,v] of Object.entries(" + json.dumps(data["tab"]) + ")) sessionStorage.setItem(k,v);")

def tab_survei(browser):
    browser.cookies()
    pages = [p for p in reversed(browser.pages) if not p.is_closed() and p.url.startswith("https://fasih-sm.bps.go.id/")]
    return next((p for p in pages if urlsplit(p.url).path == "/app/surveys"), next(iter(pages), None))

def halaman_awal(page):
    return page.url.startswith("https://fasih-sm.bps.go.id/") and urlsplit(page.url).path.rstrip("/") in ("", "/app", "/app/home", "/app/dashboard", "/app/beranda", "/app/surveys")

def gangguan(page):
    bot = page.get_by_text(PESAN_BOT).and_(page.locator(':not(#se2026_nested [data-nested-view="true"], #se2026_nested [data-nested-view="true"] *, [id^="nama_usaha_edit#"], [id^="nama_usaha_edit#"] *)'))
    return bot.filter(visible=True).first

def pasang_refresh(page, pulih=None):
    pesan = gangguan(page)
    refresh = page.get_by_role("button", name="Refresh Page", exact=True).filter(visible=True)
    def ulang():
        print("  [BOT] Hard refresh di tab yang sama.", flush=True)
        if pulih is not None:
            pulih()
        page.wait_for_timeout(1000)
        try:
            with page.expect_navigation(wait_until="domcontentloaded", timeout=TIMEOUT):
                # ponytail: reload(true) khusus Firefox/Camoufox; gunakan API bypass-cache bila pindah browser.
                page.evaluate("() => location.reload(true)")
        except Error as error:
            if page.is_closed() or re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                raise
    page.remove_locator_handler(pesan)
    page.add_locator_handler(pesan, ulang, no_wait_after=True)
    page.remove_locator_handler(refresh)
    def refresh_web():
        if refresh.evaluate_all("buttons => { const b = buttons.find(b => !b.disabled); if (!b) return false; b.click(); return true; }"):
            print("  [REFRESH] Tombol web diklik; memeriksa ulang halaman.", flush=True)
            if pulih is not None:
                pulih()
    page.add_locator_handler(refresh, refresh_web, no_wait_after=True)

def jeda(cfg, sisa):
    sisa -= 1
    if sisa == 0:
        return random.uniform(cfg["rest_min"], cfg["rest_max"]) * 60000, random.randint(cfg["batch_min"], cfg["batch_max"])
    return random.uniform(cfg["delay_min"], cfg["delay_max"]) * 1000, sisa

def hitung_mundur(page, waktu, label="Istirahat batch"):
    akhir = monotonic() + waktu / 1000
    while True:
        tersisa = max(0, akhir - monotonic())
        menit, detik = divmod(ceil(tersisa), 60)
        print(f"\r{label}: {menit:02d}:{detik:02d}", end="", flush=True)
        if tersisa == 0:
            print(" - lanjut.")
            return
        expect(page.locator("body")).to_be_visible(timeout=TIMEOUT)
        page.wait_for_timeout(min(1000, tersisa * 1000))

def ulang_langkah(page, tahap, aksi, pulih):
    while True:
        if pulih():
            raise Error("Halaman direfresh; ulangi pengisian dari awal.")
        try:
            hasil = aksi()
            if pulih():
                raise Error("Halaman direfresh; ulangi pengisian dari awal.")
            return hasil
        except (Error, AssertionError) as error:
            if pulih() or page.is_closed() or re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                raise
            if isinstance(error, AssertionError) and "timeout" not in str(error).lower():
                raise
            print("  [TUNGGU] " + tahap + " | " + str(error).splitlines()[0][:160], flush=True)
            page.wait_for_timeout(100)

def jalankan(pemulihan=False):
    root = Path(__file__).resolve().parent
    cfg = json.loads((root / "config.json").read_text(encoding="utf-8"))
    for key in ("delay", "batch", "rest"):
        low, high = cfg[key + "_min"], cfg[key + "_max"]
        if type(low) not in (int, float) or type(high) not in (int, float) or not 0 <= low <= high < float("inf"):
            raise ValueError("Rentang config tidak valid: " + key)
    if any(type(cfg[k]) is not int or cfg[k] < 1 for k in ("batch_min", "batch_max")):
        raise ValueError("Batch harus bilangan bulat positif")
    path = root / cfg["excel_file"]
    book = load_workbook(path)
    sheet = book[cfg["sheet_name"]]
    headers = [str(c.value).strip() for c in sheet[1]]
    headers = ["remarks" if h == "remark" else h for h in headers]
    assert all(headers.count(k) == 1 for k in ("id_fasih", "target_usaha", *FIELDS))
    assert all(headers.count(k) <= 1 for k in ("status", "keterangan", "catatan", "remarks"))
    rows = []
    for number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, cells))
        if str(row.get("status") or "").strip().upper() not in ("", "TODO", "RETRY"):
            continue
        try:
            if not all(isinstance(row[k], str) and row[k].strip() for k in ("id_fasih", "target_usaha")):
                raise ValueError("ID FASIH dan target usaha wajib berupa teks tidak kosong")
            if row.get("remarks") is not None and not isinstance(row["remarks"], str):
                raise ValueError("Remarks harus berupa teks")
            values = {k: angka(row[k]) for k in FIELDS}
            if any(v is not None and v != v.to_integral_value() for k, v in values.items() if k.startswith("tk_")):
                raise ValueError("Jumlah pekerja harus bilangan bulat")
        except (ValueError, InvalidOperation) as error:
            if "--cek-excel" in sys.argv:
                raise
            values = error
        rows.append((number, row, values))
    if "--cek-excel" in sys.argv or not rows:
        if not rows and "--cek-excel" not in sys.argv:
            print("Tidak ada baris kosong/TODO/RETRY untuk diproses.", flush=True)
        book.close()
        return
    status = headers.index("status") + 1 if "status" in headers else sheet.max_column + 1
    sheet.cell(1, status, "status")
    keterangan = headers.index("keterangan") + 1 if "keterangan" in headers else sheet.max_column + 1
    sheet.cell(1, keterangan, "keterangan")
    catatan = headers.index("catatan") + 1 if "catatan" in headers else sheet.max_column + 1
    sheet.cell(1, catatan, "catatan")
    with path.open("r+b"):
        pass
    fingerprint = identitas(root / ".fasih-profile" / "fingerprint.json")
    headless = False
    print(f"[MULAI] {len(rows)} baris; headless={headless}. Membuka browser dan memuat sesi tersimpan.", flush=True)
    with Camoufox(persistent_context=True, user_data_dir=str(root / ".fasih-profile"), fingerprint=fingerprint, i_know_what_im_doing=True, headless=headless, humanize=True, os="windows", locale="id-ID", timeout=TIMEOUT) as browser:
        browser.set_default_timeout(TIMEOUT)
        browser.set_default_navigation_timeout(TIMEOUT)
        dimuat_ulang = False
        def pulih_bot():
            nonlocal dimuat_ulang
            dimuat_ulang = True
        def tab_baru(tab):
            try:
                if len(browser.pages) > 2:
                    tab.close()
                elif not tab.is_closed():
                    pasang_refresh(tab, pulih_bot)
            except Error:
                if not tab.is_closed():
                    raise
        browser.on("page", tab_baru)
        for tab in browser.pages:
            pasang_refresh(tab, pulih_bot)
        page = browser.pages[0] if browser.pages else browser.new_page()
        for tab in browser.pages:
            if tab != page:
                tab.close()
        sesi(browser, page, root / ".fasih-profile" / "sesi.json")
        print("[SESI] Memeriksa akses halaman Survei...", flush=True)
        dibuka = False
        while True:
            try:
                if not dibuka:
                    dibuka = True
                    page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0&perPage=10&layout=list", wait_until="commit")
                expect(page.locator("body")).to_be_visible(timeout=TIMEOUT)
                expect(gangguan(page)).to_be_hidden(timeout=TIMEOUT)
                if not pemulihan:
                    pemulihan = True
                    input("Setelah login awal selesai, tekan Enter sekali untuk mulai otomatis... ")
                tab = tab_survei(browser)
                if tab is None:
                    page.wait_for_timeout(100)
                    continue
                page = tab
                if urlsplit(page.url).path != "/app/surveys":
                    page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0&perPage=10&layout=list", wait_until="commit")
                page.wait_for_url("**/app/surveys?*", wait_until="commit")
                sesi(browser, page, root / ".fasih-profile" / "sesi.json", simpan=True)
                break
            except (Error, AssertionError, StopIteration) as error:
                if re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                    raise
                if page.is_closed():
                    raise
                page.wait_for_timeout(100)
        print("[SESI] Halaman Survei terbuka. Mulai memeriksa assignment.", flush=True)
        hasil = dict.fromkeys(("SUCCESS", "NO_CHANGE", "ERROR"), 0)
        dimuat_ulang = False
        sisa = random.randint(cfg["batch_min"], cfg["batch_max"])
        for index, (number, row, values) in enumerate(rows):
            print(f"[{index + 1}/{len(rows)}] Baris Excel {number}: {row['target_usaha']}", flush=True)
            sheet.cell(number, status, "ERROR")
            simpan(book, path)
            assignment = page
            result, note = "ERROR", ""
            kirim_ulang = "Kirim belum selesai" in str(row.get("catatan") or "")
            tahap = "Membuka assignment detail"
            print("  " + tahap, flush=True)
            detail_dibuka = False
            dalam_edit = False
            dalam_kirim = False
            isi_ulang = False
            while True:
                try:
                    if dimuat_ulang:
                        dimuat_ulang = False
                        dalam_kirim = False
                        isi_ulang = True
                        print("  [PULIH] Mengenali ulang tampilan setelah refresh.", flush=True)
                    if isinstance(values, Exception):
                        raise values
                    if (detail_dibuka or page != assignment) and halaman_awal(page):
                        assignment = page
                        detail_dibuka = dalam_edit = dalam_kirim = False
                        isi_ulang = True
                        print("  [PULIH] Survei/Beranda terbuka; membuka kembali assignment yang sama.", flush=True)
                    if not detail_dibuka:
                        detail_dibuka = True
                        for tab in browser.pages:
                            if tab != assignment:
                                tab.close()
                        pasang_refresh(page, pulih_bot)
                        page.goto("https://fasih-sm.bps.go.id/app/assignment-detail/" + quote(row["id_fasih"], safe=""), wait_until="commit")
                    if page == assignment:
                        tahap = "Membuka tab Review"
                        print("  " + tahap, flush=True)
                        review = assignment.get_by_role("button", name="Review", exact=True)
                        expect(review).to_be_visible(timeout=TIMEOUT)
                        with assignment.expect_popup(timeout=TIMEOUT) as popup:
                            review.click(no_wait_after=True)
                        page = popup.value
                    tahap = "Memeriksa tampilan form/pengiriman"
                    print("  " + tahap, flush=True)
                    ringkasan = page.get_by_role("button", name="Ringkasan", exact=True).filter(visible=True)
                    revoke, edit = page.locator(SELECTORS["revoke"]), page.locator(SELECTORS["edit"])
                    blok = page.locator(SELECTORS["blok_ii"])
                    pengiriman = page.locator(SELECTORS["modal"]).or_(page.get_by_text("Konfirmasi Kirim", exact=True).filter(visible=True))
                    tombol_kirim = page.get_by_role("button", name="Kirim", exact=True).filter(visible=True)
                    siap = ringkasan.or_(revoke).or_(edit).or_(blok).or_(pengiriman).or_(sukses_kirim(page))
                    if dalam_kirim or kirim_ulang:
                        siap = siap.or_(tombol_kirim)
                    ulang_langkah(page, tahap, lambda: expect(siap.first).to_be_visible(timeout=1000), lambda: dimuat_ulang or halaman_awal(page))
                    if dimuat_ulang:
                        continue
                    lanjut_kirim = dalam_kirim or pengiriman.count() > 0 or (kirim_ulang and tombol_kirim.count() > 0 and not revoke.is_visible() and not edit.is_visible() and not blok.is_visible())
                    if sukses_kirim(page).is_visible():
                        result = "SUCCESS"
                    elif not lanjut_kirim:
                        dalam_edit = blok.is_visible() and not revoke.is_visible() and not edit.is_visible() and (dalam_edit or isi_ulang or not ringkasan.is_visible())
                        tahap = "Menunggu Revoke/Edit"
                        print("  " + tahap, flush=True)
                        try:
                            kontrol = revoke.or_(edit).or_(blok) if dalam_edit else revoke.or_(edit)
                            expect(kontrol.first).to_be_visible(timeout=TIMEOUT)
                        except AssertionError as error:
                            if page.is_closed() or gangguan(page).is_visible():
                                raise
                            raise ValueError("Akses edit tidak tersedia: tidak ada tombol Revoke maupun Edit.") from error
                        changed = revoke.is_visible()
                        if changed and kirim_ulang:
                            tahap = "Memeriksa hasil Kirim sebelumnya"
                            blok.click()
                            buka_usaha(page, row["target_usaha"])
                            sesuai = all(angka(inputan(page, label).input_value()) == values[column] for column, label in FIELDS.items() if values[column] is not None)
                            page.locator(SELECTORS["catatan"]).click()
                            if sesuai and page.locator(SELECTORS["switch"]).get_attribute("aria-checked") == "true":
                                result = "SUCCESS"
                        if result != "SUCCESS":
                            if changed:
                                tahap = "Revoke dan konfirmasi"
                                kirim_ulang = True
                                revoke.click()
                                page.locator(SELECTORS["konfirmasi_revoke"]).click()
                            tahap = "Membuka Edit"
                            print("  " + tahap, flush=True)
                            if not dalam_edit or changed:
                                expect(edit).to_be_visible(timeout=TIMEOUT)
                            if edit.is_visible():
                                edit.click(no_wait_after=True)
                                dalam_edit = True
                                expect(edit).to_be_hidden(timeout=TIMEOUT)
                            tahap = "Membuka BLOK II"
                            ulang_langkah(page, tahap, lambda: page.locator(SELECTORS["blok_ii"]).click(), lambda: dimuat_ulang or halaman_awal(page))
                            tahap = "BLOK II: membuka usaha " + row["target_usaha"]
                            print("  " + tahap, flush=True)
                            ulang_langkah(page, tahap, lambda: buka_usaha(page, row["target_usaha"]), lambda: dimuat_ulang or halaman_awal(page))
                            for column, label in FIELDS.items():
                                if values[column] is None:
                                    continue
                                tahap = "Mencari/mengisi " + label
                                print("  " + tahap, flush=True)
                                diubah, perlu_catatan = False, False
                                def isi_field():
                                    nonlocal kirim_ulang, changed, diubah, perlu_catatan
                                    field = inputan(page, label)
                                    try:
                                        expect(field).to_have_count(1, timeout=TIMEOUT_CARI_USAHA)
                                    except AssertionError as error:
                                        if page.is_closed() or gangguan(page).is_visible() or dimuat_ulang:
                                            raise
                                        raise ValueError("Input tidak tersedia/unik: " + label) from error
                                    berbeda = not field.input_value().strip() or angka(field.input_value()) != values[column]
                                    perlu_catatan = perlu_catatan or berbeda
                                    if berbeda or isi_ulang or diubah:
                                        diubah = kirim_ulang = changed = True
                                        field.fill(format(values[column], "f"))
                                        field.press("Tab")
                                    if angka(field.input_value()) != values[column]:
                                        raise Error("Nilai web belum sesuai Excel; ulangi kolom ini.")
                                    if perlu_catatan and (row.get("remarks") or "").strip():
                                        print("  Menyimpan remarks pada " + label, flush=True)
                                        isi_remarks(page, field, row["remarks"])
                                ulang_langkah(page, tahap, isi_field, lambda: dimuat_ulang or halaman_awal(page))
                            tahap = "CATATAN: memeriksa switch anomali"
                            print("  " + tahap, flush=True)
                            def cek_catatan():
                                nonlocal kirim_ulang, changed
                                page.locator(SELECTORS["catatan"]).click()
                                switch = page.locator(SELECTORS["switch"])
                                expect(switch).to_have_count(1, timeout=TIMEOUT)
                                if switch.get_attribute("aria-checked") == "false":
                                    kirim_ulang = changed = True
                                    page.locator(SELECTORS["kontrol_switch"]).click()
                                expect(switch).to_have_attribute("aria-checked", "true", timeout=TIMEOUT)
                            ulang_langkah(page, tahap, cek_catatan, lambda: dimuat_ulang or halaman_awal(page))
                            isi_ulang = False
                            lanjut_kirim = changed or kirim_ulang
                    if result != "SUCCESS":
                        if lanjut_kirim:
                            tahap = "Kirim: melanjutkan tampilan pengiriman"
                            print("  " + tahap, flush=True)
                            dalam_kirim = True
                            kirim_ulang = True
                            detail = kirim(page, pulih=lambda: dimuat_ulang or halaman_awal(page))
                            if dimuat_ulang:
                                continue
                            if detail:
                                cell = sheet.cell(number, keterangan, detail)
                                cell.data_type = "s"
                                note = "Galat validasi Kirim: " + detail
                            else:
                                result = "SUCCESS"
                        else:
                            result = "NO_CHANGE"
                    if result == "SUCCESS":
                        tahap = "Kirim berhasil; menutup konfirmasi sukses"
                        tutup = page.get_by_role("button", name="Tetap di Halaman", exact=True)
                        if sukses_kirim(page).is_visible():
                            tutup.click()
                    tahap = "Menyimpan sesi setelah " + result
                    sesi(browser, page, root / ".fasih-profile" / "sesi.json", simpan=True)
                    break
                except (Error, AssertionError, ValueError, InvalidOperation) as error:
                    error_text = str(error).strip()
                    if re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", error_text):
                        sheet.cell(number, status, "RETRY")
                        simpan(book, path)
                        raise
                    if result != "ERROR":
                        break
                    if halaman_awal(page):
                        continue
                    if dimuat_ulang:
                        if page.is_closed():
                            raise
                        continue
                    if isinstance(error, (ValueError, InvalidOperation)) or (isinstance(error, AssertionError) and "timeout" not in error_text.lower()):
                        note = f"{tahap} | {type(error).__name__}: {error_text}"
                        break
                    if page.is_closed():
                        raise
                    print("  [TUNGGU] " + tahap + " | " + error_text.splitlines()[0][:160], flush=True)
                    page.wait_for_timeout(100)
            for tab in browser.pages:
                if tab != assignment and not tab.is_closed() and len(browser.pages) > 1:
                    with suppress(Error):
                        tab.close()
            sheet.cell(number, status, result)
            if note:
                cell = sheet.cell(number, catatan)
                text = str(cell.value or "")
                if note not in text:
                    text += ("\n" if text else "") + note
                assert len(text) <= 32767, "Catatan melebihi batas sel Excel"
                cell.value, cell.data_type = text, "s"
            simpan(book, path)
            hasil[result] += 1
            print(f"  [HASIL] {sheet.cell(number, status).value}; progres disimpan.", flush=True)
            if note:
                print("  [CATATAN] " + note.splitlines()[0], flush=True)
            page = assignment
            page.bring_to_front()
            if index + 1 < len(rows):
                istirahat = sisa == 1
                waktu, sisa = jeda(cfg, sisa)
                if istirahat:
                    hitung_mundur(page, waktu)
                else:
                    print(f"  Jeda {waktu / 1000:.1f} detik sebelum baris berikutnya.", flush=True)
                    page.wait_for_timeout(waktu)
        else:
            print(f"[SELESAI] SUCCESS={hasil['SUCCESS']}, NO_CHANGE={hasil['NO_CHANGE']}, ERROR={hasil['ERROR']}.", flush=True)
    browser.remove_listener("page", tab_baru)
    book.close()

def main():
    jalankan()

if __name__ == "__main__":
    main()
