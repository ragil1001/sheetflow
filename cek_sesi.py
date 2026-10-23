import json
from contextlib import contextmanager, nullcontext
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Thread
from unittest import TestCase
from unittest.mock import patch

from camoufox.sync_api import Camoufox
from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error, expect

from fasih import FIELDS, PESAN_BOT, buka_usaha, identitas, inputan, jalankan, kirim, main, pasang_refresh, pesan_galat, sesi, ulang_langkah


@contextmanager
def web_lokal():
    pages = {}
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            html = pages.get(self.path, lambda: "")()
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.end_headers()
            self.wfile.write(html.encode("utf-8"))
        def log_message(self, *args):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}", pages
    finally:
        server.shutdown()
        server.server_close()
        thread.join()


with patch("fasih.jalankan") as run:
    main()
    run.assert_called_once_with()
with patch("fasih.jalankan", side_effect=Error("NS_ERROR_UNKNOWN_HOST")) as run:
    with TestCase().assertRaises(Error):
        main()
    run.assert_called_once_with()

modal = '<div class="tw:border! tw:relative tw:flex tw:w-full tw:flex-col tw:gap-4 tw:bg-background tw:p-6 tw:pt-4 tw:shadow-lg! tw:sm:rounded-lg">'
form = ('<button>Ringkasan</button><button onclick="this.hidden=true"><svg class="tabler-icon-edit" width="24" height="24"></svg></button>'
        '<div class="fasih-form-sidebar"><div title="SE2026 - L BLOK II">BLOK II</div><div title="CATATAN">CATATAN</div></div>'
        '<div id="se2026_nested"><div data-nested-view="true" onclick="this.dataset.clicked=1">USAHA CONTOH</div></div>'
        '<div id="nama_usaha_edit#1002"><input readonly value="NAMA FORM BERBEDA"></div>'
        '<label>Pekerja laki-laki<input value="0"></label>'
        '<div id="cek_anomali_button"><input role="switch" aria-checked="true"></div>')
submit = ('<script>let a=0,b=0,c=0;function gagal(stage){document.querySelector("#m").hidden=true;document.querySelector("#k").hidden=true;'
          'const first=document.querySelector("#awal");first.hidden=false;first.disabled=true;setTimeout(()=>{first.disabled=false},20);'
          'document.body.dataset[stage]=Number(document.body.dataset[stage]||0)+1}</script>'
          '<button id="awal" onclick="this.disabled=true;setTimeout(()=>{if(a++<3){gagal(\'awal\')}else{this.hidden=true;this.disabled=false;document.querySelector(\'#m\').hidden=false}},20)">Kirim</button>'
          '<section id="m" hidden>' + modal + '<button onclick="this.disabled=true;setTimeout(()=>{this.disabled=false;if(b++<2){gagal(\'modal\')}else{document.querySelector(\'#m\').hidden=true;document.querySelector(\'#k\').hidden=false}},20)">Kirim</button></div></section>'
          '<section id="k" hidden><p>Konfirmasi Kirim</p><button onclick="this.disabled=true;setTimeout(()=>{this.disabled=false;if(c++<2){gagal(\'konfirmasi\')}else{document.querySelector(\'#k\').hidden=true;document.querySelector(\'#s\').hidden=false}},20)">Konfirmasi</button></section>'
          '<section id="s" hidden><p>Assignment Berhasil DisubmitPerubahan Anda berhasil disimpan.</p><button>Tetap di Halaman</button></section>')
validasi = ('<button onclick="this.hidden=true;document.querySelector(\'#validasi\').hidden=false">Kirim</button>'
            '<section id="validasi" hidden>' + modal +
            '<button onclick="localStorage.setItem(\'galat_diklik\',\'1\');document.querySelector(\'#validasi\').hidden=true;document.querySelector(\'#rincian\').hidden=false;setTimeout(()=>{document.querySelector(\'#pesan\').hidden=false},150)">Galat 3 Perlu diperbaiki</button>'
            '<button onclick="localStorage.setItem(\'kirim_invalid\',\'1\')">Kirim</button></div></section>'
            '<section id="rincian" hidden>' + modal + '<p>Total: 0 / 3</p><section id="pesan" hidden>'
            '<div><div title="Pekerja laki-laki">Tenaga kerja</div><ul class="tw:text-destructive">'
            '<li>Jumlah pekerja harus lebih dari nol.</li><li>Jumlah pekerja tidak sesuai rincian.</li></ul></div>'
            '<div><div>Nilai pendapatan barang dan jasa</div><ul class="tw:text-destructive"><li>Pendapatan wajib diisi.</li></ul></div>'
            '</section></div></section>')
detail_galat = ('Pekerja laki-laki — Jumlah pekerja harus lebih dari nol.\n'
                'Pekerja laki-laki — Jumlah pekerja tidak sesuai rincian.\n'
                'Nilai pendapatan barang dan jasa — Pendapatan wajib diisi.')

with TemporaryDirectory() as temp, web_lokal() as (url, pages), patch("builtins.print"):
    root = Path(temp)
    fingerprint = identitas(root / "fingerprint.json")
    assert identitas(root / "fingerprint.json") == fingerprint
    cfg = dict(excel_file="uji.xlsx", sheet_name="Sheet1", delay_min=0, delay_max=0, batch_min=20, batch_max=20, rest_min=0, rest_max=0)
    (root / "config.json").write_text(json.dumps(cfg), encoding="utf-8")

    def workbook(rows, remarks=None):
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.append(["id_fasih", "target_usaha", *FIELDS, "status", "catatan", "remarks"])
        for ident, nama, value in rows:
            sheet.append([ident, nama, *[value.get(key) if isinstance(value, dict) else value if key == "tk_laki_laki" else None for key in FIELDS], "TODO", None, remarks])
        book.save(root / "uji.xlsx")
        book.close()

    with Camoufox(headless=True, os="windows", fingerprint=fingerprint, i_know_what_im_doing=True) as browser:
        context = browser.new_context()
        context.on("page", pasang_refresh)
        context.route("**/*", lambda r: r.fulfill(content_type="text/html", body="Survei"))
        page = context.new_page()
        page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0")
        page.evaluate("() => {localStorage.setItem('uji','lokal');sessionStorage.setItem('uji','lokal')}")
        context.add_cookies([dict(name="uji", value="lokal", url=page.url)])
        snapshot = root / "sesi.json"
        sesi(context, page, snapshot, simpan=True)
        assert json.loads(snapshot.read_text(encoding="utf-8"))["tab"] == {"uji": "lokal"}
        for target, kartu in (("USAHA CONTOH (PEMILIK A)", "USAHA CONTOH"), ("USAHA CONTOH", "USAHA CONTOH &lt;PEMILIK B&gt;"), ("USAHA CONTOH <PEMILIK A", "USAHA CONTOH (LAIN)")):
            page.set_content(form.replace(">USAHA CONTOH</div>", ">" + kartu + "</div>"))
            buka_usaha(page, target)
            assert page.locator('[data-clicked="1"]').count() == 1
        page.set_content(form.replace(">USAHA CONTOH</div>", ">USAHA LAIN</div>"))
        with patch("fasih.TIMEOUT_CARI_USAHA", 50), TestCase().assertRaisesRegex(ValueError, "Kartu usaha tidak cocok"):
            buka_usaha(page, "USAHA CONTOH")
        page.set_content(form.replace(">USAHA CONTOH</div>", ">USAHA BOT-12345678</div>"))
        buka_usaha(page, "USAHA BOT-12345678")
        assert not page.is_closed()
        page.set_content(submit)
        with patch("fasih.TIMEOUT", 1000):
            assert kirim(page) == ""
        assert page.evaluate("({...document.body.dataset})") == {"awal": "3", "modal": "2", "konfirmasi": "2"}
        page.set_content('<button onclick="document.body.dataset.kirim=1">Kirim</button>' + modal + '<h2>Berhasil</h2><button>Tetap di Halaman</button></div>')
        with patch.object(page, "wait_for_timeout", side_effect=AssertionError("Sukses masih ditunggu")):
            assert kirim(page) == ""
        assert page.locator("body").get_attribute("data-kirim") is None, "Sudah sukses tetapi Kirim diklik lagi"
        page.set_content('<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>')
        with patch("fasih.TIMEOUT", 50):
            assert pesan_galat(page) == "", "Modal sukses masih menunggu modal validasi"
        page.set_content('<button onclick="this.disabled=true;setTimeout(()=>{document.body.innerHTML=\'<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>\'},100)">Kirim</button>')
        with patch("fasih.TIMEOUT", 50), patch.object(page, "wait_for_timeout", wraps=page.wait_for_timeout) as tunggu:
            assert kirim(page) == ""
            assert all(c.args[0] <= 100 for c in tunggu.call_args_list), "Sukses tertunda oleh jeda tetap"
        page.set_content(modal + '<button disabled>Kirim</button></div><script>setTimeout(()=>{document.body.innerHTML=\'<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>\'},100)</script>')
        with patch.object(page, "reload", side_effect=AssertionError("Loading Kirim direfresh")):
            assert kirim(page) == "", "Kirim masih menunggu tombol pada modal lama"
        page.set_content(validasi)
        tunggu, putaran = page.wait_for_timeout, []
        def tunggu_galat(ms):
            putaran.append(ms)
            if len(putaran) > 20:
                raise RuntimeError("Rincian Galat tidak terbaca setelah tombol diklik")
            tunggu(ms)
        with patch("fasih.TIMEOUT", 50), patch.object(page, "wait_for_timeout", side_effect=tunggu_galat):
            assert kirim(page) == detail_galat
        assert pesan_galat(page) == detail_galat, "Rincian Galat yang sudah terbuka tidak dibaca ulang"
        assert page.evaluate("localStorage.getItem('galat_diklik')") == "1"
        assert page.evaluate("localStorage.getItem('kirim_invalid')") is None
        page.set_content(modal + '<button onclick="this.hidden=true;setTimeout(()=>{document.body.innerHTML=\'<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>\'},100)">Galat 1 Perlu diperbaiki</button></div>')
        assert kirim(page) == "", "Sukses masih menunggu rincian Galat lama"
        page.set_content(modal + '<button onclick="this.hidden=true;setTimeout(()=>{document.querySelector(\'#konfirmasi\').hidden=false},150)">Galat 1 Perlu diperbaiki</button>'
                         '<button id="konfirmasi" hidden onclick="document.body.innerHTML=\'<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>\'">Konfirmasi</button></div>')
        with patch("fasih.TIMEOUT", 10000), patch.object(page, "wait_for_timeout", wraps=page.wait_for_timeout) as tunggu, patch("builtins.print") as output:
            assert kirim(page) == ""
            assert len(tunggu.call_args_list) < 20, "Modal baru masih menunggu rincian Galat lama"
            assert any("Konfirmasi Kirim" in str(c) or "buka Galat" in str(c) for c in output.call_args_list)
        page.set_content(validasi.replace("document.querySelector('#validasi').hidden=true;", ""))
        assert kirim(page) == detail_galat, "Modal Galat bertumpuk membuat skrip salah memilih tombol"
        page.set_content('<button disabled onclick="document.body.innerHTML=\'<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>\'">Kirim</button><script>setTimeout(()=>document.querySelector("button").disabled=false,300)</script>')
        with patch("builtins.print") as output:
            assert kirim(page) == ""
            assert any("belum aktif/loading" in str(c) for c in output.call_args_list), "Alasan menunggu tidak ditampilkan"
        page.set_content(submit.replace("function gagal(stage){", "function gagal(stage){document.querySelector('#error').hidden=false;") + '<p id="error" hidden>Failed presigned URL</p>')
        with patch("fasih.TIMEOUT", 1000):
            assert kirim(page) == "", "Pesan presigned lama menghalangi Kirim ulang"
        for html in ("<p>Bot Detected</p>", "<p>Anda terdeteksi sebagai bot</p>", "<p>BOT-12345678</p>"):
            assert PESAN_BOT.search(html)
            requests = []
            def bot():
                requests.append(True)
                return html if len(requests) <= 3 else '<button>Ringkasan</button>'
            pages["/uji-bot"] = bot
            tab = context.new_page()
            tab.route(url + "/**", lambda r: r.continue_())
            tab.goto(url + "/uji-bot")
            tab.evaluate("sessionStorage.setItem('tetap','1')")
            tabs = list(context.pages)
            with patch.object(tab, "evaluate", wraps=tab.evaluate) as evaluate:
                expect(tab.get_by_role("button", name="Ringkasan")).to_be_visible(timeout=20000)
                assert sum(c.args == ("() => location.reload(true)",) for c in evaluate.call_args_list) == 3
            assert len(requests) == 4 and list(context.pages) == tabs
            assert tab.evaluate("sessionStorage.getItem('tetap')") == "1"
            tab.close()

        page.set_content('<p>Loading...</p>')
        page.evaluate("setTimeout(() => document.body.innerHTML='<button>Ringkasan</button>', 300)")
        with patch("fasih.TIMEOUT", 50), patch.object(page, "reload", side_effect=AssertionError("Loading direfresh")):
            expect(page.get_by_role("button", name="Ringkasan")).to_be_visible(timeout=0)
        tab = context.new_page()
        tab.set_content('<button onclick="document.body.dataset.refresh=Number(document.body.dataset.refresh||0)+1;if(document.body.dataset.refresh===\'3\')document.body.innerHTML=\'<button>Ringkasan</button>\'">Refresh Page</button>')
        with patch.object(tab, "reload", side_effect=AssertionError("Refresh browser dipakai")):
            expect(tab.get_by_role("button", name="Ringkasan")).to_be_visible(timeout=2000)
        assert not tab.is_closed() and tab.locator("body").get_attribute("data-refresh") == "3"
        tab.close()

        detail_requests = []
        def detail_assignment(route):
            ident = route.request.url.rsplit('/', 1)[-1]
            detail_requests.append(ident)
            route.fulfill(content_type="text/html", body='<button onclick="window.open(\'/form/' + ident + '\')">Review</button>')
        context.route("**/app/assignment-detail/*", detail_assignment)
        akses_requests, glitch_requests = [], {}
        def halaman(route):
            ident = route.request.url.rsplit("/", 1)[-1]
            html = form + (submit if ident == "kirim" else "")
            if ident == "kirim":
                html = html.replace('<p>Assignment Berhasil DisubmitPerubahan Anda berhasil disimpan.</p>', '<h2>Assignment Berhasil Disubmit</h2><p>Perubahan Anda berhasil disimpan.</p>')
                html = html.replace('<button>Tetap di Halaman</button>', '<button onclick="localStorage.setItem(\'sukses_ditutup\',\'1\')">Tetap di Halaman</button>')
            elif ident == "lambat":
                html = '<p>Loading...</p><template>' + form + '</template><script>setTimeout(()=>{document.body.innerHTML=document.querySelector("template").innerHTML},300)</script>'
            elif ident == "galat":
                html = form + validasi
            elif ident.startswith("putus_"):
                html = '<section id="editor">' + form + '</section>' + (validasi if ident == "putus_galat" else submit)
                html = html.replace('onclick="this.disabled=true;', 'onclick="document.querySelector(\'#editor\').hidden=true;this.disabled=true;')
                html = html.replace('onclick="this.hidden=true;', 'onclick="document.querySelector(\'#editor\').hidden=true;this.hidden=true;')
                if ident == "putus_form":
                    html = form + submit
            elif ident == "urutan":
                inputs = ''.join('<label>' + label + '<input value="0" oninput="localStorage.setItem(\'urutan\',(localStorage.getItem(\'urutan\')||\'\')+\'' + column + ',\')"></label>' for column, label in FIELDS.items())
                html = form.replace('<label>Pekerja laki-laki<input value="0"></label>', inputs) + submit.replace("a++<3", "false").replace("b++<2", "false").replace("c++<2", "false")
                html = html.replace('<div title="CATATAN">', '<div title="CATATAN" onclick="localStorage.setItem(\'urutan\',localStorage.getItem(\'urutan\')+\'CATATAN,\')">')
            elif ident.startswith("pembelian-"):
                html = form
                field = '<label>Biaya pembelian barang yang terjual<input value="9" oninput="localStorage.setItem(\'' + ident + '\',this.value)"></label>'
                if ident in ("pembelian-ada0", "pembelian-ada1"):
                    html += field + submit
                elif ident == "pembelian-lambat0":
                    html += '<section id="pembelian" hidden>' + field + '</section><script>setTimeout(()=>document.querySelector("#pembelian").hidden=false,150)</script>' + submit
                elif ident == "pembelian-duplikat0":
                    html += field * 2
            elif ident == "edit_lambat":
                html = form.replace('<button onclick="this.hidden=true">', '<button id="edit" hidden onclick="this.hidden=true;document.querySelector(\'.fasih-form-sidebar\').hidden=false;document.querySelector(\'button\').hidden=true">')
                html = html.replace('<div class="fasih-form-sidebar">', '<div class="fasih-form-sidebar" hidden>')
                html += ('<button onclick="document.querySelector(\'#revoke\').hidden=false"><svg class="tabler-icon-rotate-clockwise" width="24" height="24"></svg></button>'
                         '<div id="revoke" role="alertdialog" hidden>Revoke Assignment?<button onclick="this.parentElement.hidden=true;document.querySelector(\'.tabler-icon-rotate-clockwise\').parentElement.hidden=true;setTimeout(()=>document.querySelector(\'#edit\').hidden=false,300)">Konfirmasi</button></div>' + submit)
            elif ident.startswith("glitch-"):
                glitch_requests[ident] = glitch_requests.get(ident, 0) + 1
                html = form + submit.replace("a++<3", "false").replace("b++<2", "false").replace("c++<2", "false")
                html = html.replace('onclick="this.hidden=true"><svg class="tabler-icon-edit"', 'onclick="this.hidden=true;document.querySelectorAll(\'label,[role=switch]\').forEach(e=>e.hidden=false)"><svg class="tabler-icon-edit"')
                if ident in ("glitch-ringkasan", "glitch-ringkasan-timeout"):
                    jeda_form = 40 if ident == "glitch-ringkasan" else 350
                    html = '<button>Ringkasan</button><section id="isi" hidden>' + html + '</section><script>setTimeout(()=>document.querySelector("#isi").hidden=false,' + str(jeda_form) + ')</script>'
                elif glitch_requests[ident] == 1 and ident in ("glitch-edit", "glitch-revoke", "glitch-konfirmasi"):
                    if ident != "glitch-edit":
                        html = html.replace("tabler-icon-edit", "tabler-icon-rotate-clockwise")
                    html += '<script>document.querySelector("svg").parentElement.disabled=true</script>'
                    if ident == "glitch-konfirmasi":
                        html += '<div role="alertdialog">Revoke Assignment?<button disabled>Konfirmasi</button></div>'
            if ident == "beda":
                html = html.replace(">USAHA CONTOH</div>", ">USAHA LAIN</div>")
            elif ident == "akses":
                akses_requests.append(True)
                if len(akses_requests) == 1:
                    html = html.replace('<button onclick="this.hidden=true"><svg class="tabler-icon-edit" width="24" height="24"></svg></button>', "")
                    html = html.replace('input value="0"', 'input readonly value="0"')
                    html = html.replace('input role="switch"', 'input disabled role="switch"')
            route.fulfill(content_type="text/html", body=html)
        context.route("**/form/*", halaman)
        workbook([("beda", "USAHA CONTOH", None), ("akses", "USAHA CONTOH", None), ("angka", "USAHA CONTOH", -1), ("lanjut", "USAHA CONTOH (PEMILIK A)", None), ("lambat", "USAHA CONTOH", None), ("galat", "USAHA CONTOH", 1), ("kirim", "USAHA CONTOH", 1), ("edit_lambat", "USAHA CONTOH", 1)])
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1000), patch("fasih.TIMEOUT_CARI_USAHA", 50), \
             patch("fasih.Camoufox", side_effect=[nullcontext(context)]), patch("builtins.input", return_value=""), patch("builtins.print") as output, \
             patch("playwright.sync_api.Page.reload", side_effect=AssertionError("Akses gagal direfresh, bukan membuka Review ulang")):
            main()
            assert not any("[TUNGGU]" in str(c) and "Kirim:" not in str(c) for c in output.call_args_list), "Edit lambat menyebabkan langkah salah/timeout"
        book = load_workbook(root / "uji.xlsx")
        rows = list(book["Sheet1"].values)
        assert [row[rows[0].index("status")] for row in rows[1:]] == ["ERROR", "NO_CHANGE", "ERROR", "NO_CHANGE", "NO_CHANGE", "ERROR", "SUCCESS", "SUCCESS"], rows
        assert "Kartu usaha tidak cocok" in rows[1][rows[0].index("catatan")]
        assert len(akses_requests) == 2 and detail_requests.count("akses") == 1, "Pemulihan akses tidak membuka Review dari detail yang sama"
        assert "Angka tidak valid" in rows[3][rows[0].index("catatan")]
        assert rows[6][rows[0].index("keterangan")] == detail_galat
        assert rows[6][rows[0].index("catatan")] == "Galat validasi Kirim: " + detail_galat
        assert book["Sheet1"].cell(7, rows[0].index("keterangan") + 1).data_type == "s"
        assert context.pages[0].evaluate("localStorage.getItem('galat_diklik')") == "1"
        assert context.pages[0].evaluate("localStorage.getItem('kirim_invalid')") is None
        assert context.pages[0].evaluate("localStorage.getItem('sukses_ditutup')") == "1"
        book.close()
        context.cookies()
        assert len(context.pages) == 1
        glitches = ["ringkasan", "ringkasan-timeout", "edit", "revoke", "konfirmasi", "input", "usaha", "catatan", "akses-input", "akses-usaha", "akses-catatan", "akses-kirim"]
        workbook([("glitch-" + mode, "USAHA CONTOH", 1) for mode in glitches])
        gangguan_input, gangguan_usaha, gangguan_catatan = [], [], []
        akses_hilang = set()
        def input_glitch(tab, label):
            field = inputan(tab, label)
            if tab.url.endswith("glitch-input") and label == FIELDS["tk_laki_laki"]:
                fill = field.fill
                def isi(value):
                    if not gangguan_input:
                        gangguan_input.append(True)
                        tab.evaluate("document.querySelector('label').hidden=true;document.querySelector('.tabler-icon-edit').parentElement.hidden=false")
                        raise Error("Timeout input; tampilan kembali ke akses Edit")
                    fill(value)
                field.fill = isi
            if tab.url.endswith("glitch-akses-input") and label == FIELDS["tk_laki_laki"]:
                fill_asli = field.fill
                def isi_akses(value):
                    if "input" not in akses_hilang:
                        akses_hilang.add("input")
                        tab.set_content('<button>Ringkasan</button>')
                        raise Error("Timeout input; akses menghilang")
                    fill_asli(value)
                field.fill = isi_akses
            return field
        def usaha_glitch(tab, target):
            if tab.url.endswith("glitch-akses-usaha") and "usaha" not in akses_hilang:
                akses_hilang.add("usaha")
                tab.set_content('<button>Ringkasan</button>')
                raise ValueError("Kartu menghilang karena akses menghilang")
            if tab.url.endswith("glitch-usaha") and not gangguan_usaha:
                gangguan_usaha.append(True)
                tab.set_content(validasi)
                raise ValueError("Kartu usaha menghilang saat modal Galat terbuka")
            return buka_usaha(tab, target)
        def catatan_glitch(tab, tahap, aksi, pulih):
            def langkah():
                if tab.url.endswith("glitch-akses-catatan") and tahap.startswith("CATATAN") and "catatan" not in akses_hilang:
                    akses_hilang.add("catatan")
                    tab.set_content('<button>Ringkasan</button>')
                    raise Error("Timeout CATATAN; akses menghilang")
                if tab.url.endswith("glitch-catatan") and tahap.startswith("CATATAN") and not gangguan_catatan:
                    gangguan_catatan.append(True)
                    tab.evaluate("document.querySelector('.tabler-icon-edit').parentElement.hidden=false;document.querySelector('[role=switch]').hidden=true")
                    raise Error("Timeout CATATAN; tampilan kembali ke akses Edit")
                return aksi()
            return ulang_langkah(tab, tahap, langkah, pulih)
        def kirim_glitch(tab, pulih=None):
            if tab.url.endswith("glitch-akses-kirim") and "kirim" not in akses_hilang:
                akses_hilang.add("kirim")
                tab.set_content('<button>Ringkasan</button>')
            return kirim(tab, pulih=pulih)
        def cek_glitch(*args, **kwargs):
            if "Baris Excel" in " ".join(map(str, args)):
                cek_glitch.log.clear()
            cek_glitch.log.append(" ".join(map(str, args)))
            if len(cek_glitch.log) > 100:
                raise RuntimeError("Retry glitch tidak menentukan ulang aksi: " + "\n".join(cek_glitch.log[-12:]))
        cek_glitch.log = []
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 500), patch("fasih.TIMEOUT_CARI_USAHA", 100), \
             patch("fasih.Camoufox", return_value=nullcontext(context)), patch("builtins.input", return_value="") as masuk, \
             patch("fasih.inputan", side_effect=input_glitch), patch("fasih.buka_usaha", side_effect=usaha_glitch), \
             patch("fasih.ulang_langkah", side_effect=catatan_glitch), patch("fasih.kirim", side_effect=kirim_glitch), \
             patch("playwright.sync_api.Page.reload", side_effect=AssertionError("Pemulihan akses memakai refresh")), patch("builtins.print", side_effect=cek_glitch):
            try:
                main()
            except Error as error:
                raise RuntimeError("\n".join(cek_glitch.log[-20:])) from error
            masuk.assert_called_once()
        book = load_workbook(root / "uji.xlsx")
        rows = list(book["Sheet1"].values)
        assert [row[rows[0].index("status")] for row in rows[1:]] == ["SUCCESS"] * 6 + ["ERROR"] + ["SUCCESS"] * 5, rows
        assert rows[7][rows[0].index("keterangan")] == detail_galat
        assert glitch_requests == {"glitch-" + mode: 2 if mode in ("edit", "revoke", "konfirmasi") or mode.startswith("akses-") else 1 for mode in glitches}, glitch_requests
        assert all(detail_requests.count("glitch-" + mode) == 1 for mode in glitches), "Pemulihan akses meninggalkan detail ID FASIH semula"
        assert akses_hilang == {"input", "usaha", "catatan", "kirim"}
        assert gangguan_input == gangguan_usaha == gangguan_catatan == [True]
        book.close()
        workbook([("putus_" + tahap, "USAHA CONTOH", 1) for tahap in ("awal", "modal", "konfirmasi", "galat", "form")])
        terputus = set()
        def lanjut_kirim(tab, pulih=None):
            tahap = tab.url.rsplit("_", 1)[-1]
            if tahap not in terputus:
                terputus.add(tahap)
                if tahap == "form":
                    raise Error("Kirim terputus sementara; form edit kembali aktif")
                tujuan = {"awal": "#awal", "modal": '#m:not([hidden])', "konfirmasi": '#k:not([hidden])', "galat": '#validasi:not([hidden])'}[tahap]
                if tahap == "awal":
                    tab.locator("#awal").click()
                    tab.wait_for_timeout(50)
                else:
                    for _ in range(12):
                        if tab.locator(tujuan).is_visible():
                            break
                        tab.get_by_role("button", name="Kirim", exact=True).filter(visible=True).click()
                        tab.wait_for_timeout(50)
                    assert tab.locator(tujuan).is_visible(), "Fixture tahap Kirim tidak terbuka"
                    if tahap == "galat":
                        tab.get_by_role("button", name="Galat 3 Perlu diperbaiki", exact=True).click()
                raise Error("Kirim terputus sementara; tampilan pengiriman masih terbuka")
            return kirim(tab, pulih=pulih)
        def cek_tunggu(*args, **kwargs):
            text = " ".join(map(str, args))
            if "[TUNGGU]" in text and "Kirim:" not in text:
                raise RuntimeError("Tampilan Kirim aktif tetapi skrip kembali menunggu tahap form")
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 500), patch("fasih.Camoufox", return_value=nullcontext(context)), \
             patch("builtins.input", return_value=""), patch("fasih.kirim", side_effect=lanjut_kirim), patch("builtins.print", side_effect=cek_tunggu):
            main()
        book = load_workbook(root / "uji.xlsx")
        rows = list(book["Sheet1"].values)
        assert terputus == {"awal", "modal", "konfirmasi", "galat", "form"}
        assert [row[rows[0].index("status")] for row in rows[1:]] == ["SUCCESS", "SUCCESS", "SUCCESS", "ERROR", "SUCCESS"], rows
        assert rows[4][rows[0].index("keterangan")] == detail_galat
        book.close()
        workbook([("urutan", "USAHA CONTOH", dict.fromkeys(FIELDS, 1))])
        gagal_input = []
        def input_terputus(tab, label):
            field = inputan(tab, label)
            if label == FIELDS["perbaikan_biaya_pembelian"]:
                tekan_asli = field.press
                def tekan(key):
                    if not gagal_input:
                        gagal_input.append(True)
                        raise Error("Timeout sesudah mengisi biaya pembelian")
                    tekan_asli(key)
                field.press = tekan
            return field
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1000), patch("fasih.Camoufox", return_value=nullcontext(context)), \
             patch("builtins.input", return_value=""), patch("fasih.inputan", side_effect=input_terputus), patch("builtins.print") as output:
            main()
            assert sum("Memeriksa tampilan form/pengiriman" in str(c) for c in output.call_args_list) == 2, "Gangguan satu input mengulang pemeriksaan form"
        urutan = list(FIELDS)
        urutan.insert(7, "perbaikan_biaya_pembelian")
        assert context.pages[0].evaluate("localStorage.getItem('urutan')") == ",".join([*urutan, "CATATAN", ""])
        assert gagal_input == [True]
        book = load_workbook(root / "uji.xlsx")
        assert book["Sheet1"].cell(2, len(FIELDS) + 3).value == "SUCCESS"
        book.close()
        workbook([("lanjut", "USAHA CONTOH", None)])
        context.route("https://fasih-sm.bps.go.id/app/home", lambda r: r.fulfill(content_type="text/html", body="Beranda"))
        pemeriksaan_login = []
        def selesai_login(c):
            pemeriksaan_login.append(True)
            if len(pemeriksaan_login) < 3:
                return None
            if len(pemeriksaan_login) == 3:
                c.pages[0].goto("https://fasih-sm.bps.go.id/app/home")
            return c.pages[0]
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1000), patch("fasih.Camoufox", return_value=nullcontext(context)), \
             patch("fasih.tab_survei", side_effect=selesai_login), patch("builtins.input", return_value="") as masuk:
            main()
            masuk.assert_called_once()
        assert len(pemeriksaan_login) == 3
        pembelian = [("kosong", None), ("hilang0", 0), ("hilang1", 5), ("ada0", 0), ("ada1", 5), ("lambat0", 0), ("duplikat0", 0)]
        workbook([("pembelian-" + ident, "USAHA CONTOH", {"perbaikan_biaya_pembelian": value}) for ident, value in pembelian], remarks="Catatan uji")
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1000), patch("fasih.TIMEOUT_CARI_USAHA", 300), \
             patch("fasih.Camoufox", return_value=nullcontext(context)), patch("builtins.input", return_value=""), patch("fasih.isi_remarks") as catatan:
            main()
            assert catatan.call_count == 3, "Input tidak ada/kosong tetap diberi remarks"
        book = load_workbook(root / "uji.xlsx")
        rows = list(book["Sheet1"].values)
        assert [row[rows[0].index("status")] for row in rows[1:]] == ["NO_CHANGE", "NO_CHANGE", "ERROR", "SUCCESS", "SUCCESS", "SUCCESS", "ERROR"], rows
        assert "Input tidak tersedia/unik: Biaya pembelian" in rows[3][rows[0].index("catatan")]
        assert "Input tidak tersedia/unik: Biaya pembelian" in rows[7][rows[0].index("catatan")]
        assert context.pages[0].evaluate("localStorage.getItem('pembelian-ada0')") == "0"
        assert context.pages[0].evaluate("localStorage.getItem('pembelian-ada1')") == "5"
        assert context.pages[0].evaluate("localStorage.getItem('pembelian-lambat0')") == "0"
        book.close()
        pulih_requests = {}
        def halaman_pulih(ident):
            pulih_requests[ident] = pulih_requests.get(ident, 0) + 1
            jumlah = pulih_requests[ident]
            if jumlah == 2:
                return "<p>Bot Detected</p>"
            pemicu, mode = ident.split("_")
            if jumlah == 3 and mode in ("home", "surveys"):
                return '<script>location.replace("https://fasih-sm.bps.go.id/app/' + mode + '")</script>'
            if jumlah == 3 and mode == "detail":
                return '<button onclick="window.open(\'/form/pulih/' + ident + '\')">Review</button>'
            assert jumlah <= (4 if mode in ("home", "surveys", "detail") else 3), "Form terus direfresh setelah BOT hilang"
            if jumlah == 3 and mode in ("awal", "modal", "konfirmasi", "galat", "sukses"):
                if mode == "sukses":
                    return '<h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button>'
                if mode == "galat":
                    return validasi
                return submit + '<script>' + {"awal": "", "modal": 'document.querySelector("#awal").hidden=true;document.querySelector("#m").hidden=false;', "konfirmasi": 'document.querySelector("#awal").hidden=true;document.querySelector("#k").hidden=false;'}[mode] + '</script>'
            html = form.replace('input value="0"', 'input id="laki" value="1"' if jumlah >= 3 else 'input id="laki" value="0"')
            html += '<label>Pekerja perempuan<input id="perempuan" value="0"></label>'
            html = html.replace('<input role="switch" aria-checked="true">', '<input role="switch" aria-checked="false"><div id="uji-control" onclick="this.previousElementSibling.setAttribute(\'aria-checked\',\'true\')">Aktifkan</div>')
            if jumlah == 3 and mode == "edit":
                html = html.replace('<button onclick="this.hidden=true"><svg class="tabler-icon-edit" width="24" height="24"></svg></button>', "")
            if jumlah == 3 and mode == "revoke":
                html = html.replace('<button onclick="this.hidden=true">', '<button id="edit" hidden onclick="this.hidden=true">')
                html += '<div role="alertdialog">Revoke Assignment?<button onclick="this.parentElement.hidden=true;document.querySelector(\'#edit\').hidden=false">Konfirmasi</button></div>'
            html += ('<button id="awal" onclick="kirimTahap(\'awal\')">Kirim</button><section id="m" hidden>' + modal +
                     '<button onclick="kirimTahap(\'modal\')">Kirim</button></div></section>'
                     '<section id="k" hidden><p>Konfirmasi Kirim</p><button onclick="kirimTahap(\'konfirmasi\')">Konfirmasi</button></section>'
                     '<section id="s" hidden><h2>Assignment Berhasil Disubmit</h2><button>Tetap di Halaman</button></section>'
                     '<script>const diisi={};document.querySelectorAll("#laki,#perempuan").forEach(input=>input.addEventListener("input",()=>{diisi[input.id]=true;'
                     + ('if(input.id==="perempuan")document.body.innerHTML="<p>Bot Detected</p>";' if jumlah == 1 and pemicu == "input" else '') + '}));'
                     'function kirimTahap(tahap){'
                     + ('if(tahap===' + json.dumps(pemicu) + '){document.body.innerHTML="<p>Bot Detected</p>";return;}' if jumlah == 1 else '') +
                     'if(tahap==="awal"){document.querySelector("#awal").hidden=true;document.querySelector("#m").hidden=false;}'
                     'else if(tahap==="modal"){document.querySelector("#m").hidden=true;document.querySelector("#k").hidden=false;}'
                     'else{localStorage.setItem(' + json.dumps(ident) + ',JSON.stringify({diisi,laki:document.querySelector("#laki").value,perempuan:document.querySelector("#perempuan").value,switch:document.querySelector("[role=switch]").getAttribute("aria-checked")}));document.querySelector("#k").hidden=true;document.querySelector("#s").hidden=false;}}</script>')
            return html
        ids_pulih = [pemicu + "_" + mode for pemicu in ("input", "awal", "modal", "konfirmasi") for mode in ("review", "edit")] + ["konfirmasi_" + mode for mode in ("home", "surveys", "detail", "revoke", "awal", "modal", "konfirmasi", "galat", "sukses")]
        for ident in ids_pulih:
            pages["/form/pulih/" + ident] = lambda ident=ident: halaman_pulih(ident)
        context.route("**/app/assignment-detail/pulih-*", lambda r: r.fulfill(content_type="text/html", body='<button onclick="window.open(\'' + url + '/form/pulih/' + r.request.url.rsplit('pulih-', 1)[-1] + '\')">Review</button>'))
        context.route("https://fasih-sm.bps.go.id/app/home", lambda r: r.fulfill(content_type="text/html", body="Beranda"))
        workbook([("pulih-" + ident, "USAHA CONTOH", {"tk_laki_laki": 1, "tk_perempuan": 2}) for ident in ids_pulih])
        context.route(url + "/**", lambda r: r.continue_())
        def cek_pulih(*args, **kwargs):
            if "[TUNGGU]" in " ".join(map(str, args)) and "Memeriksa tampilan" in " ".join(map(str, args)):
                raise RuntimeError("BOT sudah hilang tetapi tahap form tidak dikenali")
        with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1500), patch("fasih.Camoufox", return_value=nullcontext(context)), \
             patch("builtins.input", return_value="") as masuk, patch("builtins.print", side_effect=cek_pulih):
            main()
            masuk.assert_called_once()
        book = load_workbook(root / "uji.xlsx")
        rows = list(book["Sheet1"].values)
        assert [row[rows[0].index("status")] for row in rows[1:]] == ["ERROR" if ident.endswith("_galat") else "SUCCESS" for ident in ids_pulih], rows
        assert rows[ids_pulih.index("konfirmasi_galat") + 1][rows[0].index("keterangan")] == detail_galat
        assert pulih_requests == {ident: 4 if ident.endswith(("home", "surveys", "detail")) else 3 for ident in ids_pulih}
        tersimpan = {item["name"]: item["value"] for origin in context.storage_state()["origins"] if origin["origin"] == url for item in origin["localStorage"]}
        for ident in ids_pulih:
            if ident.rsplit("_", 1)[-1] in ("awal", "modal", "konfirmasi", "galat", "sukses"):
                continue
            data = json.loads(tersimpan[ident])
            assert data == {"diisi": {"laki": True, "perempuan": True}, "laki": "1", "perempuan": "2", "switch": "true"}, (ident, data)
        book.close()
        diterima = []
        def terkirim():
            if diterima:
                html = form.replace("tabler-icon-edit", "tabler-icon-rotate-clockwise").replace('input value="0"', 'input readonly value="1"')
            else:
                html = (form + '<button onclick="this.hidden=true;document.querySelector(\'#m\').hidden=false">Kirim</button>'
                        '<section id="m" hidden>' + modal + '<button onclick="document.querySelector(\'#m\').hidden=true;document.querySelector(\'#k\').hidden=false">Kirim</button></div></section>'
                        '<section id="k" hidden><p>Konfirmasi Kirim</p><button onclick="fetch(\'/terima\').then(()=>{document.body.innerHTML=\'<p>Bot Detected</p>\'})">Konfirmasi</button></section>')
            if pemulihan == "akses":
                return html.replace("<p>Bot Detected</p>", '<button>Ringkasan</button>')
            return html.replace("<p>Bot Detected</p>", '<button onclick=location.reload()>Refresh Page</button>') if pemulihan == "refresh" else html
        pages["/form/diterima"] = terkirim
        def terima_lokal():
            diterima.append(True)
            return "ok"
        pages["/terima"] = terima_lokal
        context.route(url + "/**", lambda r: r.continue_())
        context.route("**/app/assignment-detail/diterima", lambda r: r.fulfill(content_type="text/html", body='<button onclick="window.open(\'' + url + '/form/diterima\')">Review</button>'))
        for pemulihan in ("bot", "refresh", "akses"):
            diterima.clear()
            workbook([("diterima", "USAHA CONTOH", {"tk_laki_laki": 1, "perbaikan_biaya_pembelian": 0})])
            with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1000), patch("fasih.TIMEOUT_CARI_USAHA", 50), patch("fasih.Camoufox", return_value=nullcontext(context)), patch("builtins.input", return_value=""):
                assert jalankan() is None
            assert diterima == [True], "Kiriman yang sudah diterima dikirim ulang"
            book = load_workbook(root / "uji.xlsx")
            assert book["Sheet1"].cell(2, len(FIELDS) + 3).value == "SUCCESS"
            book.close()
        context.close()

    workbook([("pulih", "USAHA CONTOH", None)])
    launches, closed, surveys = [], [], []
    @contextmanager
    def buka(**options):
        assert not launches, "BOT membuka ulang browser"
        with Camoufox(**{**options, "headless": True, "timeout": 90000}) as context:
            launches.append(options["user_data_dir"])
            context.route("**/*", lambda r: r.fulfill(content_type="text/html", body="Survei"))
            def survei():
                surveys.append(True)
                return "Bot Detected" if len(surveys) < 5 else "Survei"
            pages["/app/surveys?page=0&perPage=10&layout=list"] = survei
            context.route(url + "/**", lambda r: r.continue_())
            context.route("https://fasih-sm.bps.go.id/app/surveys?*", lambda r: r.fulfill(status=302, headers={"location": url + "/app/surveys?page=0&perPage=10&layout=list"}))
            context.route("**/app/assignment-detail/*", lambda r: r.fulfill(content_type="text/html", body='<button onclick="window.open(\'/form\')">Review</button>'))
            context.route("**/form", lambda r: r.fulfill(content_type="text/html", body=form))
            yield context
        closed.append(True)
        assert not context.pages
    with patch("fasih.__file__", str(root / "fasih.py")), patch("fasih.TIMEOUT", 1500), patch("fasih.Camoufox", side_effect=buka), \
         patch("fasih.tab_survei", side_effect=lambda c: c.pages[0]), \
         patch("builtins.input", return_value=""):
        main()
        assert len(launches) == len(closed) == 1 and len(surveys) == 5
    book = load_workbook(root / "uji.xlsx")
    assert book["Sheet1"].cell(2, len(FIELDS) + 3).value == "NO_CHANGE"
    book.close()

print("Pemeriksaan glitch/posisi aktif, buka ulang Review saat akses hilang, biaya pembelian/remarks, urutan input, Galat/Excel, dan pemulihan BOT lulus.")
