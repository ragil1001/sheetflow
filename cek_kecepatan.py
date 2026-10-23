"""Bandingkan alur lengkap dengan HEAD memakai workbook/profil sementara dan web fixture."""
import argparse
import json
import subprocess
from collections import defaultdict
from contextlib import contextmanager
from html import escape
from pathlib import Path
from tempfile import TemporaryDirectory
from time import perf_counter
from types import ModuleType
from unittest.mock import patch

from camoufox.sync_api import Camoufox
from openpyxl import Workbook, load_workbook

import fasih


def ukur(module, root):
    (root / "config.json").write_text(json.dumps(dict(excel_file="uji.xlsx", sheet_name="Sheet1", delay_min=0,
        delay_max=0, batch_min=20, batch_max=20, rest_min=0, rest_max=0)), encoding="utf-8")
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["id_fasih", "target_usaha", *module.FIELDS, "status", "remarks"])
    sheet.append(["uji", "USAHA UJI", *[1] * len(module.FIELDS), "TODO", "Catatan uji"])
    book.save(root / "uji.xlsx")
    book.close()
    fields = "".join('<div id="field#' + column + '" class="tw:rounded-lg tw:gap-2"><p>' + escape(label) +
        '</p><input value="0" oninput="urutan.push(\'' + column + '\')">'
        '<button class="tw:size-8 tw:relative tw:text-primary tw:@lg/form:flex! tw:hidden" '
        'onclick="document.querySelector(\'#remarks\').hidden=false"><svg class="tw:size-4" width="24" height="24"></svg>Catatan</button></div>'
        for column, label in module.FIELDS.items())
    html = ('<button>Ringkasan</button><button onclick="this.hidden=true"><svg class="tabler-icon-edit" width="24" height="24"></svg>Edit</button>'
        '<div class="fasih-form-sidebar"><div title="SE2026 - L BLOK II">BLOK II</div><div title="CATATAN" onclick="urutan.push(\'CATATAN\')">CATATAN</div></div>'
        '<div id="se2026_nested"><div data-nested-view="true">USAHA UJI</div></div>' + fields +
        ''.join('<div class="tw:rounded-lg tw:gap-2"><p>Pertanyaan lain ' + str(i) + '</p><input readonly value="0"></div>' for i in range(200)) +
        '<div id="cek_anomali_button"><input role="switch" aria-checked="true"></div>'
        '<section id="remarks" hidden><textarea placeholder="Tambah catatan di sini..."></textarea><button onclick="urutan.push(\'remarks\')">Simpan</button>'
        '<button onclick="this.parentElement.hidden=true">Dismiss</button></section>'
        '<button onclick="urutan.push(\'Kirim\');this.hidden=true;document.querySelector(\'#validasi\').hidden=false">Kirim</button>'
        '<section id="validasi" hidden><div class="tw:border! tw:relative tw:flex tw:w-full tw:flex-col tw:gap-4 tw:bg-background tw:p-6 tw:pt-4 tw:shadow-lg! tw:sm:rounded-lg">'
        '<button onclick="urutan.push(\'Kirim modal\');document.querySelector(\'#validasi\').hidden=true;document.querySelector(\'#konfirmasi\').hidden=false">Kirim</button></div></section>'
        '<section id="konfirmasi" hidden><p>Konfirmasi Kirim</p><button onclick="localStorage.setItem(\'urutan\',JSON.stringify(urutan));'
        'this.parentElement.hidden=true;document.querySelector(\'#sukses\').hidden=false">Konfirmasi</button></section>'
        '<section id="sukses" hidden><p>Assignment Berhasil Disubmit</p><button>Tetap di Halaman</button></section><script>const urutan=[]</script>')
    def route(r):
        if "/app/assignment-detail/" in r.request.url:
            body = '<button onclick="window.open(\'/form\')">Review</button>'
        else:
            body = html if r.request.url.endswith("/form") else "Survei"
        r.fulfill(content_type="text/html", body=body)
    @contextmanager
    def buka(**options):
        with Camoufox(**{**options, "headless": True}) as context:
            context.route("**/*", route)
            page = context.new_page()
            page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0")
            start = perf_counter()
            yield context
            elapsed.append(perf_counter() - start)
            assert json.loads(context.pages[0].evaluate("localStorage.getItem('urutan')")) == expected
    elapsed = []
    expected = [item for column in module.FIELDS for item in (column, "remarks")] + ["CATATAN", "Kirim", "Kirim modal"]
    totals, counts = defaultdict(float), defaultdict(int)
    def timer(name, original):
        def wrapped(*args, **kwargs):
            start = perf_counter()
            try:
                return original(*args, **kwargs)
            finally:
                totals[name] += perf_counter() - start
                counts[name] += 1
        return wrapped
    with patch.object(module, "__file__", str(root / "fasih.py")), patch.object(module, "Camoufox", side_effect=buka), \
         patch("builtins.input", return_value=""), patch("builtins.print"), \
         patch.object(module, "posisi", side_effect=timer("posisi", module.posisi)), \
         patch.object(module, "ulang_langkah", side_effect=timer("langkah", module.ulang_langkah)), \
         patch.object(module, "isi_remarks", side_effect=timer("remarks", module.isi_remarks)), \
         patch.object(module, "kirim", side_effect=timer("kirim", module.kirim)), \
         patch.object(module, "sesi", side_effect=timer("sesi", module.sesi)), \
         patch.object(module, "simpan", side_effect=timer("excel", module.simpan)):
        module.main()
    book = load_workbook(root / "uji.xlsx")
    assert book["Sheet1"].cell(2, len(module.FIELDS) + 3).value == "SUCCESS"
    book.close()
    print(f"Total {elapsed[0]:.2f}s; " + "; ".join(f"{key}={totals[key]:.2f}s/{counts[key]}x" for key in totals))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", nargs="?", const="HEAD", help="Ukur ref Git sebagai pembanding (default HEAD)")
    args = parser.parse_args()
    module = fasih
    if args.baseline:
        result = subprocess.run(["git", "show", args.baseline + ":fasih.py"], cwd=Path(__file__).parent, capture_output=True, text=True, encoding="utf-8", check=True)
        module = ModuleType("fasih_baseline")
        module.__file__ = fasih.__file__
        exec(compile(result.stdout, "fasih_baseline.py", "exec"), module.__dict__)
    with TemporaryDirectory() as temp:
        print("Baseline " + args.baseline if args.baseline else "Lokal")
        ukur(module, Path(temp))
